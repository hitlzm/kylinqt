from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SKILL_SCRIPTS = Path(
    r"C:\Users\hitlz\.codex\plugins\cache\openai-primary-runtime"
    r"\documents\26.1025.11240\skills\documents\scripts"
)
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import (  # noqa: E402
    apply_table_geometry,
    column_widths_from_weights,
    section_content_width_dxa,
)


BLACK = RGBColor(0x00, 0x00, 0x00)
TEXT = RGBColor(0x20, 0x24, 0x2A)
MUTED = RGBColor(0x5B, 0x63, 0x6D)
NAVY = "1F3A5F"
PALE_BLUE = "F4F7FB"
BORDER = "D9D9D9"
FONT_CN = "Microsoft YaHei"
FONT_CODE = "Consolas"


def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None,
                 color: RGBColor | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_style(style, name: str, size: float, bold: bool = False,
              color: RGBColor = BLACK) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    rpr.rFonts.set(qn("w:eastAsia"), FONT_CN)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = BORDER, size: int = 4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def clear_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def clear_style_paragraph_borders(style) -> None:
    p_pr = style.element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def keep_table_rows_together(table) -> None:
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = tr_pr.find(qn("w:cantSplit"))
        if cant_split is None:
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)


def repeat_header(table) -> None:
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def add_page_number_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    run = paragraph.add_run("外引导跟踪模式  |  第 ")
    set_run_font(run, FONT_CN, 8.5, False, MUTED)

    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)
    set_run_font(run, FONT_CN, 8.5, False, MUTED)

    run = paragraph.add_run(" 页")
    set_run_font(run, FONT_CN, 8.5, False, MUTED)


def add_rich_paragraph(container, text: str, *, size: float = 10.5,
                       color: RGBColor = TEXT, bold_lead: bool = False,
                       space_after: float = 6, space_before: float = 0,
                       first_line_indent: float | None = None,
                       alignment=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph = container.add_paragraph()
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.15
    if first_line_indent is not None:
        paragraph.paragraph_format.first_line_indent = Inches(first_line_indent)

    tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    lead_pending = bold_lead
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, FONT_CN, size, True, color)
            lead_pending = False
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, FONT_CODE, size - 0.7, False, color)
            lead_pending = False
        else:
            run = paragraph.add_run(token)
            set_run_font(run, FONT_CN, size, lead_pending, color)
            lead_pending = False
    return paragraph


def add_heading(document, text: str, level: int = 1,
                *, page_break_before: bool = False) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break_before
    run = paragraph.add_run(text)
    set_run_font(
        run,
        FONT_CN,
        15 if level == 1 else 12,
        True,
        BLACK,
    )


def add_caption(document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run, FONT_CN, 9.5, True, BLACK)


def add_data_table(document, caption: str, headers: list[str], rows: list[list[str]],
                   weights: list[float], *, font_size: float = 8.8,
                   alignments: list[WD_ALIGN_PARAGRAPH] | None = None) -> None:
    add_caption(document, caption)
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = 1

    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        run = paragraph.add_run(text)
        set_run_font(run, FONT_CN, font_size, True, RGBColor(0xFF, 0xFF, 0xFF))
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "FFFFFF" if row_index % 2 == 0 else PALE_BLUE)
            set_cell_borders(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                alignments[col_index]
                if alignments and col_index < len(alignments)
                else (WD_ALIGN_PARAGRAPH.CENTER if col_index == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            )
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            text_parts = re.split(r"(`.*?`)", value)
            for part in text_parts:
                if not part:
                    continue
                line_parts = part.split("\n")
                if part.startswith("`") and part.endswith("`"):
                    first_line = line_parts[0]
                    if len(line_parts) == 1:
                        first_line = first_line[1:-1]
                    else:
                        first_line = first_line[1:]
                    run = paragraph.add_run(first_line)
                    set_run_font(run, FONT_CODE, font_size - 0.4, False, TEXT)
                else:
                    run = paragraph.add_run(line_parts[0])
                    set_run_font(
                        run,
                        FONT_CN,
                        font_size,
                        col_index == 0,
                        TEXT,
                    )
                for line in line_parts[1:]:
                    if line.endswith("`"):
                        line = line[:-1]
                    run.add_break()
                    if line:
                        run = paragraph.add_run(line)
                        set_run_font(
                            run,
                            FONT_CODE if part.startswith("`") else FONT_CN,
                            font_size - 0.4 if part.startswith("`") else font_size,
                            col_index == 0,
                            TEXT,
                        )

    width_dxa = section_content_width_dxa(document.sections[0])
    widths = column_widths_from_weights(weights, width_dxa)
    apply_table_geometry(table, widths, cell_margins_dxa={
        "top": 105,
        "bottom": 105,
        "start": 130,
        "end": 130,
    })
    repeat_header(table)
    keep_table_rows_together(table)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.34)

    styles = document.styles
    set_style(styles["Normal"], FONT_CN, 10.5, False, TEXT)
    normal = styles["Normal"].paragraph_format
    normal.space_before = Pt(0)
    normal.space_after = Pt(6)
    normal.line_spacing = 1.15
    normal.widow_control = True

    set_style(styles["Title"], FONT_CN, 24, True, BLACK)
    clear_style_paragraph_borders(styles["Title"])
    title = styles["Title"].paragraph_format
    title.space_before = Pt(8)
    title.space_after = Pt(8)
    title.keep_with_next = True

    set_style(styles["Heading 1"], FONT_CN, 15, True, BLACK)
    h1 = styles["Heading 1"].paragraph_format
    h1.space_before = Pt(12)
    h1.space_after = Pt(6)
    h1.keep_with_next = True

    set_style(styles["Heading 2"], FONT_CN, 12, True, BLACK)
    h2 = styles["Heading 2"].paragraph_format
    h2.space_before = Pt(8)
    h2.space_after = Pt(4)
    h2.keep_with_next = True

    add_page_number_footer(section)

    core = document.core_properties
    core.title = "外引导跟踪模式两种不同周期说明"
    core.subject = "1s 与 20ms 外引导跟踪模式的实现和时间设置"
    core.author = "kylinqt project"
    core.keywords = "外引导, 1s, 20ms, 跟踪周期, 转台, 时间同步"
    core.comments = "基于 final4 分支提交 3dae055 的代码整理"


def build_document(output_path: Path) -> None:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("外引导跟踪模式两种不同周期说明")
    set_run_font(run, FONT_CN, 24, True, BLACK)
    clear_paragraph_borders(title)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(0)
    meta.paragraph_format.space_after = Pt(12)
    run = meta.add_run("代码基线  final4 / 3dae055    整理日期  2026-10-27")
    set_run_font(run, FONT_CN, 8.8, False, MUTED)

    add_rich_paragraph(
        doc,
        "本文根据当前上位机代码，梳理外引导模式下 1s 与 20ms 两种跟踪周期的实现。"
        "两种模式共用外引导源选择、串口工作线程、转台时间同步和角度变换，但发送节拍、"
        "轨迹生成方式以及转台时间字段不同。理解这些差异后，才能判断某一模式为什么会出现"
        "预测误差、采样陈旧或端到端延迟。",
        size=10.8,
        space_after=8,
    )
    add_rich_paragraph(
        doc,
        "**核心区别是输出时间粒度。** 1s 模式每秒下发一组包含 4 个角度的轨迹，"
        "组内点间隔 250 ms，并通过 Alpha-Beta 滤波器预测未来角度；20ms 模式约每 20 ms "
        "下发一个单点联合帧，直接使用导引头最近一次解算角度。1s 模式降低串口和转台指令"
        "数量，20ms 模式提高下发给定频率，但不会改变导引头数据本身的采集、图像处理或"
        "识别延迟。",
        size=10.8,
        space_after=10,
    )

    add_heading(doc, "两种模式的核心参数")
    add_data_table(
        doc,
        "表 1  1s 与 20ms 外引导模式对照",
        ["比较项", "1s 跟踪模式", "20ms 跟踪模式"],
        [
            ["界面设置码", "7", "6"],
            ["源线程定时器", "Qt PreciseTimer，单次定时并续排", "Qt PreciseTimer，20 ms 重复定时"],
            ["发送节拍", "每秒 1 组；首次约 800 ms 后发送", "每约 20 ms 发送 1 帧；首次在定时器触发后发送"],
            ["每组内容", "4 个预测角，步长 250 ms", "1 个当前角，三轴联合帧"],
            ["转台命令", "0x0A，外框与中框各一帧", "0x0B，内框、中框、外框一帧"],
            ["时间字段", "整秒 trackTime", "20 ms 对齐的毫秒 timeMs"],
            ["角度来源", "Alpha-Beta 滤波后的外推值", "最近一次源端解算角度"],
            ["时间同步", "绑定模式时发 0x10，秒值清零", "绑定模式时发 0x10，秒值清零"],
        ],
        [1.05, 1.55, 1.55],
        font_size=8.9,
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "一、公共控制路径")
    add_heading(doc, "1. 周期设置如何进入工作线程", level=2)
    add_rich_paragraph(
        doc,
        "外引导页面将 1s 按钮映射为设置值 7，将 20ms 按钮映射为设置值 6。"
        "`ModeController` 保存运行模式、外引导源和周期，并通过 `modeChanged`、"
        "`exguideSrcChanged`、`exguideSettingChanged` 三类信号通知图像导引头、激光导引头、"
        "CCD 相机和转台工作线程。各导引头线程只在“外引导模式 + 本路来源 + 有效周期”"
        "同时成立时启动自己的发送定时器。",
        size=10.5,
    )
    add_rich_paragraph(
        doc,
        "切换来源或周期时，旧定时器先停止并解绑，旧周期状态被清除，滤波器标记为需要重新"
        "初始化。这样可避免沿用上一路导引头或上一模式留下的角速度。退出外引导时，来源和"
        "周期也会被清空，再次进入时必须重新选择。",
        size=10.5,
    )

    add_heading(doc, "2. 三层时间关系", level=2)
    add_data_table(
        doc,
        "表 2  外引导链路中的时间层级",
        ["层级", "1s 模式", "20ms 模式"],
        [
            ["源数据更新", "图像约 20 ms；激光约 10 ms；CCD 按检测帧", "与 1s 模式相同，取决于来源"],
            ["轨迹计算", "每个新测量更新 Alpha-Beta 状态", "跳过预测，直接取最新测量角"],
            ["下发调度", "本地单调时钟，每秒一组，提前 200 ms", "20 ms 软件节拍，按转台时基对齐"],
            ["转台目标时基", "转台状态帧时间，取下一整秒", "转台状态帧时间，取下一个 20 ms 边界"],
        ],
        [0.85, 1.75, 1.75],
        font_size=8.9,
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "二、1s 跟踪模式的实现")
    add_heading(doc, "1. 本地发送节奏", level=2)
    add_rich_paragraph(
        doc,
        "1s 模式使用 `PeriodicExGuideSchedule` 维护单调时间轴。基本周期为 1000 ms，"
        "发送提前量为 200 ms，因此第 k 组数据的本地发送时刻近似为 "
        "`t(k) = 1000k - 200 ms`。外引导开始后，第一组在约 800 ms 时发送，"
        "后续每组相隔约 1 s。定时器设置为单次触发；每次发送完成后，调度器根据已经过时间"
        "重新计算到下一次发送的延时，而不是简单重新启动 1000 ms 定时器。",
        size=10.5,
    )
    add_rich_paragraph(
        doc,
        "这个续排方式有两个作用。第一，回调晚到时下一次发送仍能向原绝对节拍收敛，"
        "不会把一个小的线程抖动永久累积到后续每个周期。第二，如果一次回调跨过了目标秒"
        "边界，调度器会跳到下一个目标秒，避免补发已经过期的组，但也会跳过本次已错过的槽位。",
        size=10.5,
    )

    add_heading(doc, "2. 一组数据内的 250 ms 粒度", level=2)
    add_rich_paragraph(
        doc,
        "1s 模式每秒只发一组，但一组并不是一个终点。协议命令 0x0A 内部仍带有 4 个角度点，"
        "相对整秒目标分别对应 0.25 s、0.50 s、0.75 s 和 1.00 s。发送前的 200 ms 提前量"
        "通过 `firstPredictionDelaySec` 平移到预测窗口：发送时刻越早，滤波器当前状态到"
        "第一个角度的外推时间越长。",
        size=10.5,
    )
    add_data_table(
        doc,
        "表 3  1s 模式的时序示例",
        ["本地时刻", "目标秒", "四点在目标秒内的位置", "含义"],
        [
            ["0.80 s", "1 s", "1.25 / 1.50 / 1.75 / 2.00 s", "提前 200 ms 发送目标秒 1 的整组轨迹"],
            ["1.80 s", "2 s", "2.25 / 2.50 / 2.75 / 3.00 s", "下一组继续覆盖目标秒 2 之后的 1 s"],
        ],
        [0.72, 0.62, 2.25, 2.25],
        font_size=8.9,
        alignments=[
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
    )
    add_rich_paragraph(
        doc,
        "因此，“1s 周期”描述的是上位机下发一组的频率；“250 ms”描述的是组内轨迹的"
        "时间分辨率。协议代码中的 `CMD_TRACK_250` 名称来自后者，并不表示界面选择的 1s "
        "周期写错。",
        size=10.5,
    )

    add_heading(doc, "3. 预测角如何生成", level=2)
    add_rich_paragraph(
        doc,
        "各导引头在工作线程中把最新测量角送入 `ABTrackManager`。图像导引头收到一帧后"
        "将虚拟时间戳增加 20 ms，激光导引头增加 10 ms；CCD 使用当前系统毫秒时间。"
        "Alpha-Beta 跟踪器默认参数为 alpha 0.50、beta 0.15，角速度上限 180 °/s，"
        "并按导引头视场设置角度限幅与创新门限。周期或来源切换后的首帧会重新初始化"
        "角度状态和角速度，防止旧速度在目标已经停住时继续推动首个预测包。",
        size=10.5,
    )
    add_rich_paragraph(
        doc,
        "发送时，滤波器以当前状态为起点，分别预测四个未来时刻，生成方位角和俯仰角各一组"
        "`sendExGuideData`。转台线程再进行导引头画面零位补偿、绕内框轴的反滚转变换，"
        "并把相对角叠加到当前中框、外框角度上。1s 模式发送外框 0x0A 和中框 0x0A 两帧，"
        "内框不发送跟踪指令，保持状态反馈中的当前角度。",
        size=10.5,
    )

    add_heading(doc, "4. 时间同步与最终时间字段", level=2)
    add_rich_paragraph(
        doc,
        "开始 1s 外引导绑定时，源线程发出 `reqTimesync(0)`，转台线程发送 0x10 指令并把"
        "转台秒时基设置为 0。随后转台把旧时间反馈标记为无效，等待下一帧状态反馈刷新时基。"
        "上位机调度器产生的 `frame.time` 最终不会被直接采用；转台发送函数会根据最新状态帧"
        "估算当前转台时间，并取下一整秒作为 0x0A 的 `trackTime`。",
        size=10.5,
    )
    add_rich_paragraph(
        doc,
        "这种做法的重点是让转台接收数据的时基与转台自身反馈一致，避免 0x10 丢失或被拒绝后"
        "仍使用上位机本地时间。代价是：预测窗口仍由上位机单调调度器计算，因此该模式依赖"
        "一次时间清零后两侧节拍足够接近。若 0x10 丢失或被拒绝，上位机会在下一帧状态反馈后"
        "继续发送，并使用转台原有时间；此时两侧可能产生相位偏差，需要通过现场时间记录确认。",
        size=10.5,
    )

    add_heading(doc, "三、20ms 跟踪模式的实现")
    add_heading(doc, "1. 发送节拍与数据直通", level=2)
    add_rich_paragraph(
        doc,
        "20ms 模式使用非单次 Qt PreciseTimer，`start(20)` 后按约 20 ms 软件节拍重复触发，"
        "名义频率约为 50 Hz。每次触发直接把当前的 `m_azimuth` 和 `m_pitch` 发给转台线程，"
        "不调用 Alpha-Beta 的 `GenAxisPacket`，也不生成四个未来点。",
        size=10.5,
    )
    add_rich_paragraph(
        doc,
        "因此，20ms 是“下发节拍”，不等同于“每个 20 ms 必定出现一份全新的导引头测量”。"
        "图像源约 20 ms 一帧，激光源更快，CCD 则取决于视频帧率与检测推理速度。如果新测量"
        "晚于定时器，代码会对最新值再发送一次；如果新测量早于定时器，较新的值会在下一个"
        "20 ms 节拍被取走。这是“最新值直通”的做法，而不是把每个源数据帧与每个下发帧一一配对。",
        size=10.5,
    )

    add_heading(doc, "2. 转台时间如何对齐", level=2)
    add_rich_paragraph(
        doc,
        "20ms 模式发送前，转台线程使用最近一帧状态反馈估算当前转台时间："
        "`估算时间 = 最近反馈时间 + 距反馈的已过时间 + 2 ms`，结果按 14,400,000 ms "
        "的四小时循环取模。代码允许状态帧最长陈旧 100 ms；超过该值就暂不发送，缺少有效"
        "状态帧时也不会用本地时间代替。",
        size=10.5,
    )
    add_rich_paragraph(
        doc,
        "目标时刻取估算时间之后的第一个 20 ms 边界，即下一毫秒数为 0、20、40、… 的时刻。"
        "这样一帧可以在发送后立即落入比当前时间稍晚的偶数毫秒时间点，不再固定从某个起始"
        "偏移起步。代码还保存上一帧目标时刻，若定时器抖动使两个回调落在同一个 20 ms 周期，"
        "后一个帧会被丢弃，避免重复目标时间。",
        size=10.5,
    )

    add_heading(doc, "3. 联合帧内容", level=2)
    add_rich_paragraph(
        doc,
        "转台线程先对输入的方位角、俯仰角做与 1s 模式相同的零位和反滚转变换，再形成一帧"
        "0x0B 联合指令：内框保持当前反馈角，中框为当前角加俯仰修正，外框为当前角加方位修正。"
        "如果中框或外框结果超出允许范围，本帧不发送并输出范围提示。",
        size=10.5,
    )
    add_rich_paragraph(
        doc,
        "与 1s 模式相比，20ms 模式不需要内框以外的多轴分组命令，因此一帧即可覆盖三轴；"
        "而 1s 模式使用 0x0A 时必须分别发送外框和中框，并让内框保持不动。",
        size=10.5,
    )

    add_heading(doc, "四、两种模式的详细时间设置对照")
    add_data_table(
        doc,
        "表 4  时间基准、首次发送与目标时刻",
        ["项目", "1s 跟踪模式", "20ms 跟踪模式"],
        [
            ["调度时钟", "`QElapsedTimer` 单调时间", "最新转台状态帧时间加本机经过时间"],
            ["基本步长", "1000 ms", "20 ms"],
            ["提前量", "目标秒开始前 200 ms", "没有固定提前量；目标取下一个 20 ms 边界"],
            ["首次发送", "重启调度后约 800 ms", "0x10 发出后，约首个 20 ms 定时器触发"],
            ["转台 target", "`(当前转台秒 + 1) mod 14400`", "`(floor(t / 20) + 1) × 20 ms`"],
            ["时间范围", "0 到 14399 s，循环 4 h", "0 到 14399999 ms，循环 4 h"],
            ["时间守卫", "状态帧超过 100 ms 未更新则暂不发送", "同样超过 100 ms 则暂不发送"],
            ["重复抑制", "仅单次定时器；错槽后跳到下一目标秒", "同一个 20 ms target 只允许发送一次"],
        ],
        [0.9, 1.8, 1.8],
        font_size=8.8,
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "五、通信中断与周期切换")
    add_heading(doc, "1. 图像导引头的帧编号检查", level=2)
    add_rich_paragraph(
        doc,
        "外引导过程中，图像源线程在每次发送节拍比较帧编号。帧编号超过 1000 ms 未变化，"
        "就停止本路定时器、清除周期状态、取消图像导引头来源，并保持外引导运行模式不变。"
        "这与串口打开期间的常规通信检查不同：常规检查每 500 ms 比较一次，帧编号连续 2000 ms "
        "不变时只提示通信故障，不直接终止外引导。",
        size=10.5,
    )
    add_heading(doc, "2. 其他输入约束", level=2)
    add_rich_paragraph(
        doc,
        "图像串口未打开、激光串口未打开或 CCD 视频流未接入时，不会启动该路外引导，也不会"
        "持续输出跟踪日志。串口已打开但转台状态帧无效，或计算角度超出中框、外框范围时，"
        "转台线程同样不会下发当前跟踪帧。上述检查针对的是发送前状态，不能替代现场对转台"
        "响应、串口排队和画面处理延迟的测量。",
        size=10.5,
    )

    add_heading(doc, "六、实现与代码位置")
    add_data_table(
        doc,
        "表 5  关键实现索引",
        ["代码位置", "关键实现", "时间含义"],
        [
            ["`components/ExGuide.qml`", "1s 设置 7，20ms 设置 6；修正按钮只在 1s 模式显示", "把界面选择送入周期控制器"],
            ["`ModeControl/ModeController.h`", "保存来源与周期；切换来源或周期清旧选择", "模式与周期状态源"],
            ["`serialport/\nserialport_image.cpp`", "图像源定时器、滤波器喂数、帧编号检查", "图像约 20 ms 更新；外引导断链判定 1000 ms"],
            ["`serialport/\nserialport_laser.cpp`", "激光源定时器与 Alpha-Beta 喂数", "激光虚拟时间步长 10 ms"],
            ["`serialport/\nserialport_CCD.cpp`", "CCD 定时器与当前角度直通", "按检测结果更新时刻"],
            ["`serialport/\nexguide_timesync.h`", "1000 ms 周期、200 ms 提前量、四点预测窗", "1s 调度基准"],
            ["`kalman/AlphaBetaTracker.*`", "双轴滤波、alpha/beta、限幅、预测", "生成 1s 模式的四点轨迹"],
            ["`serialport/\nserialport_turntable_HEX.cpp`", "状态时间估算、0x10、0x0A、0x0B、20 ms 对齐", "最终目标时间与协议编码"],
            ["`main.cpp`", "周期信号到三个导引头线程和转台线程的连接", "跨线程事件边界"],
        ],
        [1.72, 2.2, 1.68],
        font_size=7.9,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "七、实施时的判断要点")
    add_rich_paragraph(
        doc,
        "从控制含义看，1s 模式适合以预测覆盖 1 s 目标段，并用 4 个 250 ms 点提高段内"
        "轨迹分辨率；其风险主要来自目标机动与预测误差。20ms 模式适合直接跟踪最新测量角，"
        "其风险主要来自源帧延迟、图像识别延迟、软件调度抖动和串口排队。两种模式都没有"
        "把导引头处理时间放进各自的 1 s 或 20 ms 常量中。",
        size=10.5,
    )
    add_rich_paragraph(
        doc,
        "`Qt::PreciseTimer` 只表示 Qt 会尽量使用较高精度定时器，并不提供硬实时保证。"
        "线程事件循环阻塞、系统负载、串口写入和反馈解析仍可能使回调晚到。因此在评估实际"
        "性能时，应同时记录源帧时间、导引头线程发包时间、转台线程实际写入时间和转台反馈"
        "时间，不能只看 QML 上的 1s 或 20ms 标签。",
        size=10.5,
    )
    add_rich_paragraph(
        doc,
        "**验证时重点看三件事：** 第一，1s 组间隔是否稳定在约 1000 ms，四点间隔是否保持"
        "250 ms 的语义；第二，20ms 帧的 `timeMs` 是否持续递增且落在 20 ms 网格上；第三，"
        "源帧变慢时，两种模式是否按设计分别表现为预测继续发送或最新值重复发送。",
        size=10.5,
        space_before=3,
        space_after=0,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


if __name__ == "__main__":
    target = Path(
        sys.argv[1]
        if len(sys.argv) > 1
        else r"E:\QTproject\kylinqt\.codex-preview\tracking_modes.docx"
    )
    build_document(target)
    print(target)
